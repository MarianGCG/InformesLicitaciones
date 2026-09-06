from io import BytesIO
from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def obtener_fuente(tamano, negrita=False):
    """
    Busca una fuente disponible en Windows o Linux.
    """

    rutas = []

    if negrita:
        rutas.extend([
            r"C:\Windows\Fonts\arialbd.ttf",
            "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        ])
    else:
        rutas.extend([
            r"C:\Windows\Fonts\arial.ttf",
            "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        ])

    for ruta in rutas:
        try:
            return ImageFont.truetype(ruta, tamano)
        except Exception:
            pass

    return ImageFont.load_default()


def texto(valor):
    """
    Convierte cualquier valor a texto limpio.
    """
    if valor is None:
        return "-"

    valor = str(valor).strip()

    return valor if valor else "-"


def formatear_fecha(fecha):
    if not fecha:
        return "-"

    return fecha.strftime("%d/%m/%Y")


def obtener_empresa_registro(registro):
    """
    Devuelve la Empresa asociada al registro.
    """

    if registro.empresa_oferente:
        return registro.empresa_oferente

    if registro.empresa_proveedor:
        return registro.empresa_proveedor

    return None

def agregar_tabla_resultados(documento, registros, empresa):
    """
    Agrega una tabla REAL de Word con los resultados.
    """

    columnas = [
        "Empresa",
        "OC",
        "Comprador",
        "Proceso",
        "R",
        "Empresas que también licitaron",
        "Apertura",
        "Estado",
        "Oferta",
        "Mon",
        "Descripción",
    ]

    tabla = documento.add_table(
        rows=1,
        cols=len(columnas)
    )

    tabla.style = "Table Grid"

    # --------------------------------------------------------
    # ENCABEZADOS
    # --------------------------------------------------------

    encabezados = tabla.rows[0].cells

    for indice, nombre_columna in enumerate(columnas):
        encabezados[indice].text = nombre_columna

        for parrafo in encabezados[indice].paragraphs:
            for run in parrafo.runs:
                run.bold = True
                run.font.size = Pt(8)

    # --------------------------------------------------------
    # FILAS
    # --------------------------------------------------------

    for registro in registros:

        fila = tabla.add_row().cells

        if registro.numero_oc:
            oc = registro.numero_oc
        else:
            oc = "-"

        oferta = (
            f"{registro.precio_total_oferta:,.2f}"
            if registro.precio_total_oferta is not None
            else "-"
        )

        valores = [
            texto(empresa.nombre),
            texto(oc),
            texto(registro.comprador),
            texto(registro.numero_proceso),
            texto(registro.numero_renglon),
            texto(getattr(registro, "competidores", "")),
            formatear_fecha(registro.fecha_apertura),
            texto(registro.estado),
            oferta,
            texto(registro.moneda_oferta),
            texto(registro.descripcion_renglon),
        ]

        for indice, valor in enumerate(valores):

            fila[indice].text = str(valor)

            for parrafo in fila[indice].paragraphs:
                for run in parrafo.runs:
                    run.font.size = Pt(7)

    return tabla


def agregar_linea(documento, etiqueta, valor):
    """
    Agrega una línea:
    Etiqueta: valor
    """

    parrafo = documento.add_paragraph()

    run = parrafo.add_run(
        f"{etiqueta}: "
    )

    run.bold = True

    parrafo.add_run(
        texto(valor)
    )

    return parrafo


def generar_word_mails(empresas_con_registros):
    """
    Genera un único Word.

    empresas_con_registros:
        [
            (empresa, [registro1, registro2, ...]),
            ...
        ]
    """

    documento = Document()

    # --------------------------------------------------------
    # CONFIGURACIÓN DE PÁGINA
    # --------------------------------------------------------

    seccion = documento.sections[0]

    seccion.top_margin = Inches(0.55)
    seccion.bottom_margin = Inches(0.55)
    seccion.left_margin = Inches(0.55)
    seccion.right_margin = Inches(0.55)

    for indice, (
        empresa,
        registros
    ) in enumerate(empresas_con_registros):

        # ----------------------------------------------------
        # NUEVA PÁGINA
        # ----------------------------------------------------

        if indice > 0:
            documento.add_page_break()

        # ----------------------------------------------------
        # DESTINATARIO
        # ----------------------------------------------------

        nombre_pila = (
            empresa.nombre_pila
            or empresa.nombre
            or ""
        ).strip()

        tipo = (
            empresa.tipo_destinatario
            or "empresa"
        ).lower()

        if tipo == "persona":

            destinatario = (
                f"Hola {nombre_pila}, "
                f"¿Cómo estás?"
            )

            texto_mail = (
                f"Hola {nombre_pila}, ¿Cómo estás?\n\n"
                "Quería acercarte algo concreto: "
                "preparamos una simulación de cotización "
                "de las Cauciones de Mantenimiento de Oferta "
                "(MO) de algunos procesos que identificamos "
                "a partir de información pública, para que "
                "puedas conocer cuánto te hubiéramos cotizado "
                "nosotros.\n\n"
                "Tenemos dos alternativas:\n\n"
                "1. Te cotizamos la MO para mostrarte cuál "
                "hubiera sido nuestro costo.\n\n"
                "2. Si me enviás la póliza que contrataste "
                "oportunamente, podemos comparar y evaluar "
                "la posibilidad de mejorar el costo hasta un 30%.\n\n"
                "Sujeto, naturalmente, a las condiciones de "
                "emisión y evaluación de la compañía.\n\n"
                "La idea es que tengas una referencia concreta "
                "para próximas licitaciones.\n\n"
                "Saludos,\n"
                "Andrés"
            )

        else:

            destinatario = (
                f"Estimados {nombre_pila},"
            )

            texto_mail = (
                f"Estimados {nombre_pila},\n\n"
                "Quería acercarles algo concreto: "
                "preparamos una simulación de cotización "
                "de las Cauciones de Mantenimiento de Oferta "
                "(MO) de algunos procesos que identificamos "
                "a partir de información pública, para que "
                "puedan conocer cuánto les hubiéramos cotizado "
                "nosotros.\n\n"
                "Tenemos dos alternativas:\n\n"
                "1. Les cotizamos la MO para mostrarles cuál "
                "hubiera sido nuestro costo.\n\n"
                "2. Si nos envían la póliza que contrataron "
                "oportunamente, podemos comparar y evaluar "
                "la posibilidad de mejorar el costo hasta un 30%.\n\n"
                "Sujeto, naturalmente, a las condiciones de "
                "emisión y evaluación de la compañía.\n\n"
                "La idea es que tengan una referencia concreta "
                "para próximas licitaciones.\n\n"
                "Saludos,\n"
                "Andrés"
            )

        # ----------------------------------------------------
        # DESTINATARIO
        # ----------------------------------------------------

        agregar_linea(
            documento,
            "Destinatario",
            destinatario
        )

        # ----------------------------------------------------
        # ASUNTO
        # ----------------------------------------------------

        agregar_linea(
            documento,
            "Asunto",
            "Simulación de cotización de Cauciones de MO"
        )

        # ----------------------------------------------------
        # TEXTO
        # ----------------------------------------------------

        parrafo = documento.add_paragraph()

        run = parrafo.add_run(
            "Texto del mail:"
        )

        run.bold = True

        # ----------------------------------------------------
        # CUERPO DEL MAIL
        # ----------------------------------------------------

        for linea in texto_mail.split("\n"):

            parrafo = documento.add_paragraph()

            parrafo.paragraph_format.space_after = Pt(0)

            run = parrafo.add_run(linea)

            run.font.size = Pt(10)

        # ----------------------------------------------------
        # RESULTADOS
        # ----------------------------------------------------

        documento.add_paragraph()

        parrafo = documento.add_paragraph()

        run = parrafo.add_run(
            "Resultados:"
        )

        run.bold = True

        agregar_tabla_resultados(
            documento,
            registros,
            empresa
        )
    # --------------------------------------------------------
    # GUARDAR
    # --------------------------------------------------------

    buffer = BytesIO()

    documento.save(buffer)

    buffer.seek(0)

    return buffer